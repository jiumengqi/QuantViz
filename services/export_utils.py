# -*- coding: utf-8 -*-
"""
数据导出工具模块
支持导出为CSV、Excel和PDF格式
"""

import os
import csv
import io
import logging
from datetime import datetime

logger = logging.getLogger(__name__)


def export_to_csv(data, filename, columns=None):
    """
    导出数据为CSV格式

    Args:
        data: 数据列表，每项为字典或列表
        filename: 保存文件名（不含路径）
        columns: 可选，指定列顺序

    Returns:
        str: 完整的文件保存路径
    """
    if not data:
        raise ValueError("数据为空，无法导出")

    # 确定输出目录
    output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'exports')
    os.makedirs(output_dir, exist_ok=True)

    # 构建完整文件路径
    filepath = os.path.join(output_dir, filename)
    if not filepath.endswith('.csv'):
        filepath += '.csv'

    # 如果data是字典列表
    if isinstance(data, list) and len(data) > 0 and isinstance(data[0], dict):
        if columns is None:
            columns = list(data[0].keys())

        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.DictWriter(f, fieldnames=columns, extrasaction='ignore')
            writer.writeheader()
            writer.writerows(data)
    else:
        # data是简单列表
        with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            if columns:
                writer.writerow(columns)
            for row in data:
                if isinstance(row, (list, tuple)):
                    writer.writerow(row)
                else:
                    writer.writerow([row])

    logger.info(f"CSV导出成功: {filepath}")
    return filepath


def export_to_csv_response(data, filename, columns=None):
    """
    导出数据为CSV格式，用于Flask响应

    Args:
        data: 数据列表，每项为字典
        filename: 下载文件名（不含路径）
        columns: 可选，指定列顺序

    Returns:
        tuple: (response, filename) 其中response是Flask响应对象
    """
    from flask import Response

    if not data:
        raise ValueError("数据为空，无法导出")

    # 确定列顺序
    if columns is None and isinstance(data[0], dict):
        columns = list(data[0].keys())

    # 创建CSV内容
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=columns, extrasaction='ignore')
    writer.writeheader()
    writer.writerows(data)

    # 创建Flask响应
    response = Response(
        output.getvalue(),
        mimetype='text/csv;charset=utf-8-sig',
        headers={'Content-Disposition': f'attachment; filename*=UTF-8\'\'{filename}'}
    )

    return response, filename


def export_to_excel(data, filename, sheet_name='Sheet1', columns=None):
    """
    导出数据为Excel格式（使用openpyxl）

    Args:
        data: 数据列表，每项为字典
        filename: 保存文件名（不含路径）
        sheet_name: 工作表名称
        columns: 可选，指定列顺序

    Returns:
        str: 完整的文件保存路径
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        logger.error("openpyxl未安装，请运行: pip install openpyxl")
        raise ImportError("请先安装openpyxl: pip install openpyxl")

    if not data:
        raise ValueError("数据为空，无法导出")

    # 确定列顺序
    if columns is None and isinstance(data[0], dict):
        columns = list(data[0].keys())

    # 确定输出目录
    output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'exports')
    os.makedirs(output_dir, exist_ok=True)

    # 构建完整文件路径
    filepath = os.path.join(output_dir, filename)
    if not filepath.endswith('.xlsx'):
        filepath += '.xlsx'

    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 设置样式
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 写入表头
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 写入数据
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row_data.get(col_name, '') if isinstance(row_data, dict) else row_data[col_idx - 1] if col_idx <= len(row_data) else ''
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # 自动调整列宽
    for col_idx, col_name in enumerate(columns, 1):
        max_length = len(str(col_name))
        for row_idx in range(2, len(data) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)].width = min(max_length + 2, 50)

    # 保存文件
    wb.save(filepath)
    logger.info(f"Excel导出成功: {filepath}")
    return filepath


def export_to_excel_response(data, filename, sheet_name='Sheet1', columns=None):
    """
    导出数据为Excel格式，用于Flask响应

    Args:
        data: 数据列表，每项为字典
        filename: 下载文件名（不含路径）
        sheet_name: 工作表名称
        columns: 可选，指定列顺序

    Returns:
        tuple: (response, filename) 其中response是Flask响应对象
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
    except ImportError:
        logger.error("openpyxl未安装，请运行: pip install openpyxl")
        raise ImportError("请先安装openpyxl: pip install openpyxl")

    if not data:
        raise ValueError("数据为空，无法导出")

    # 确定列顺序
    if columns is None and isinstance(data[0], dict):
        columns = list(data[0].keys())

    # 创建工作簿
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name

    # 设置样式
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center')
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # 写入表头
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # 写入数据
    for row_idx, row_data in enumerate(data, 2):
        for col_idx, col_name in enumerate(columns, 1):
            value = row_data.get(col_name, '') if isinstance(row_data, dict) else row_data[col_idx - 1] if col_idx <= len(row_data) else ''
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(horizontal='center', vertical='center')

    # 自动调整列宽
    for col_idx, col_name in enumerate(columns, 1):
        max_length = len(str(col_name))
        for row_idx in range(2, len(data) + 2):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_length = max(max_length, len(str(cell_value)))
        ws.column_dimensions[chr(64 + col_idx) if col_idx <= 26 else 'A' + chr(64 + col_idx - 26)].width = min(max_length + 2, 50)

    # 保存到字节流
    from io import BytesIO
    output = BytesIO()
    wb.save(output)
    output.seek(0)

    from flask import Response
    response = Response(
        output.getvalue(),
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        headers={'Content-Disposition': f'attachment; filename*=UTF-8\'\'{filename}'}
    )

    return response, filename


def export_to_pdf(html_content, filename):
    """
    导出HTML内容为PDF格式（使用weasyprint）

    Args:
        html_content: HTML格式的内容
        filename: 保存文件名（不含路径）

    Returns:
        str: 完整的文件保存路径
    """
    try:
        from weasyprint import HTML
    except ImportError:
        logger.error("weasyprint未安装，请运行: pip install weasyprint")
        raise ImportError("请先安装weasyprint: pip install weasyprint")

    if not html_content:
        raise ValueError("HTML内容为空，无法导出")

    # 确定输出目录
    output_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'exports')
    os.makedirs(output_dir, exist_ok=True)

    # 构建完整文件路径
    filepath = os.path.join(output_dir, filename)
    if not filepath.endswith('.pdf'):
        filepath += '.pdf'

    # 生成PDF
    HTML(string=html_content).write_pdf(filepath)
    logger.info(f"PDF导出成功: {filepath}")
    return filepath


def export_to_pdf_response(html_content, filename):
    """
    导出HTML内容为PDF格式，用于Flask响应（使用weasyprint）

    Args:
        html_content: HTML格式的内容
        filename: 下载文件名（不含路径）

    Returns:
        tuple: (response, filename) 其中response是Flask响应对象
    """
    try:
        from weasyprint import HTML
    except ImportError:
        logger.error("weasyprint未安装，请运行: pip install weasyprint")
        raise ImportError("请先安装weasyprint: pip install weasyprint")

    if not html_content:
        raise ValueError("HTML内容为空，无法导出")

    # 确保文件名有.pdf扩展名
    if not filename.endswith('.pdf'):
        filename += '.pdf'

    # 生成PDF到字节流
    from io import BytesIO
    output = BytesIO()
    HTML(string=html_content).write_pdf(output)
    output.seek(0)

    from flask import Response
    response = Response(
        output.getvalue(),
        mimetype='application/pdf',
        headers={'Content-Disposition': f'attachment; filename*=UTF-8\'\'{filename}'}
    )

    return response, filename


def generate_backtest_report_html(result, trades_data=None, equity_data=None):
    """
    生成回测报告的HTML内容

    Args:
        result: 回测结果字典
        trades_data: 交易记录列表
        equity_data: 净值曲线数据

    Returns:
        str: HTML格式的回测报告
    """
    from datetime import datetime

    # 格式化日期
    start_date = result.get('start_date', '')
    end_date = result.get('end_date', '')
    if hasattr(start_date, 'strftime'):
        start_date = start_date.strftime('%Y-%m-%d')
    if hasattr(end_date, 'strftime'):
        end_date = end_date.strftime('%Y-%m-%d')

    # 生成交易记录HTML
    trades_html = ''
    if trades_data:
        trades_html = '''
        <table class="table">
            <thead>
                <tr>
                    <th>日期</th>
                    <th>类型</th>
                    <th>价格</th>
                    <th>数量</th>
                    <th>金额</th>
                    <th>手续费</th>
                </tr>
            </thead>
            <tbody>
        '''
        for trade in trades_data:
            trades_html += f'''
                <tr>
                    <td>{trade.get('date', '')}</td>
                    <td>{trade.get('type', '')}</td>
                    <td>{trade.get('price', 0):.2f}</td>
                    <td>{trade.get('shares', 0)}</td>
                    <td>{trade.get('amount', 0):.2f}</td>
                    <td>{trade.get('fee', 0):.2f}</td>
                </tr>
            '''
        trades_html += '</tbody></table>'

    html = f'''
    <!DOCTYPE html>
    <html lang="zh-CN">
    <head>
        <meta charset="UTF-8">
        <title>回测报告 - {result.get('strategy_name', '策略')}</title>
        <style>
            body {{
                font-family: "Microsoft YaHei", Arial, sans-serif;
                margin: 40px;
                color: #333;
            }}
            h1 {{
                color: #1a1a1a;
                border-bottom: 2px solid #366092;
                padding-bottom: 10px;
            }}
            h2 {{
                color: #366092;
                margin-top: 30px;
            }}
            .header {{
                margin-bottom: 30px;
            }}
            .info-grid {{
                display: grid;
                grid-template-columns: repeat(3, 1fr);
                gap: 15px;
                margin: 20px 0;
            }}
            .info-item {{
                background: #f5f5f5;
                padding: 15px;
                border-radius: 5px;
            }}
            .info-label {{
                font-size: 12px;
                color: #666;
            }}
            .info-value {{
                font-size: 18px;
                font-weight: bold;
                color: #333;
            }}
            .metrics-grid {{
                display: grid;
                grid-template-columns: repeat(3, 1fr);
                gap: 20px;
                margin: 20px 0;
            }}
            .metric-card {{
                background: #f8f9fa;
                border: 1px solid #ddd;
                border-radius: 8px;
                padding: 20px;
                text-align: center;
            }}
            .metric-value {{
                font-size: 24px;
                font-weight: bold;
                color: #366092;
            }}
            .metric-label {{
                font-size: 14px;
                color: #666;
                margin-top: 5px;
            }}
            .positive {{
                color: #28a745;
            }}
            .negative {{
                color: #dc3545;
            }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin-top: 20px;
            }}
            th, td {{
                border: 1px solid #ddd;
                padding: 10px;
                text-align: left;
            }}
            th {{
                background: #366092;
                color: white;
            }}
            tr:nth-child(even) {{
                background: #f9f9f9;
            }}
            .footer {{
                margin-top: 40px;
                text-align: center;
                color: #666;
                font-size: 12px;
            }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>策略回测报告</h1>
            <p><strong>策略名称：</strong>{result.get('strategy_name', 'N/A')}</p>
            <p><strong>股票代码：</strong>{result.get('stock_code', 'N/A')}</p>
            <p><strong>回测区间：</strong>{start_date} 至 {end_date}</p>
            <p><strong>生成时间：</strong>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </div>

        <h2>关键指标</h2>
        <div class="metrics-grid">
            <div class="metric-card">
                <div class="metric-value {'positive' if result.get('total_return', 0) >= 0 else 'negative'}">{result.get('total_return', 0):.2f}%</div>
                <div class="metric-label">总收益率</div>
            </div>
            <div class="metric-card">
                <div class="metric-value {'positive' if result.get('annual_return', 0) >= 0 else 'negative'}">{result.get('annual_return', 0):.2f}%</div>
                <div class="metric-label">年化收益率</div>
            </div>
            <div class="metric-card">
                <div class="metric-value">{result.get('sharpe_ratio', 0):.2f}</div>
                <div class="metric-label">夏普比率</div>
            </div>
            <div class="metric-card">
                <div class="metric-value negative">{result.get('max_drawdown', 0):.2f}%</div>
                <div class="metric-label">最大回撤</div>
            </div>
            <div class="metric-card">
                <div class="metric-value">{result.get('win_rate', 0):.2f}%</div>
                <div class="metric-label">胜率</div>
            </div>
            <div class="metric-card">
                <div class="metric-value">{result.get('total_trades', 0)}</div>
                <div class="metric-label">交易次数</div>
            </div>
        </div>

        <h2>回测参数</h2>
        <div class="info-grid">
            <div class="info-item">
                <div class="info-label">初始资金</div>
                <div class="info-value">¥{result.get('initial_capital', 0):,.2f}</div>
            </div>
            <div class="info-item">
                <div class="info-label">最终资产</div>
                <div class="info-value">¥{result.get('final_assets', 0):,.2f}</div>
            </div>
            <div class="info-item">
                <div class="info-label">盈利金额</div>
                <div class="info-value {'positive' if (result.get('final_assets', 0) - result.get('initial_capital', 0)) >= 0 else 'negative'}">
                    ¥{result.get('final_assets', 0) - result.get('initial_capital', 0):,.2f}
                </div>
            </div>
        </div>

        <h2>交易记录</h2>
        {trades_html if trades_html else '<p>暂无交易记录</p>'}

        <div class="footer">
            <p>本报告由股票分析与量化投资平台自动生成</p>
            <p>仅供参考，不构成投资建议</p>
        </div>
    </body>
    </html>
    '''
    return html


def generate_portfolio_report_html(portfolio_data, holdings_data=None):
    """
    生成投资组合报告的HTML内容

    Args:
        portfolio_data: 投资组合信息字典
        holdings_data: 持仓数据列表

    Returns:
        str: HTML格式的投资组合报告
    """
    from datetime import datetime

    # 生成持仓HTML
    holdings_html = ''
    if holdings_data:
        holdings_html = '''
        <table class="table">
            <thead>
                <tr>
                    <th>代码</th>
                    <th>名称</th>
                    <th>数量</th>
                    <th>成本价</th>
                    <th>当前价</th>
                    <th>市值</th>
                    <th>占比</th>
                </tr>
            </thead>
            <tbody>
        '''
        for holding in holdings_data:
            holdings_html += f'''
                <tr>
                    <td>{holding.get('symbol', '')}</td>
                    <td>{holding.get('name', '')}</td>
                    <td>{holding.get('quantity', 0)}</td>
                    <td>{holding.get('avg_cost', 0):.2f}</td>
                    <td>{holding.get('current_price', 0):.2f}</td>
                    <td>{holding.get('market_value', 0):.2f}</td>
                    <td>{holding.get('allocation', 0):.2f}%</td>
                </tr>
            '''
        holdings_html += '</tbody></table>'

    html = f'''
    <!DOCTYPE html>
    <html lang="zh-CN">
    <head>
        <meta charset="UTF-8">
        <title>投资组合报告 - {portfolio_data.get('name', '投资组合')}</title>
        <style>
            body {{
                font-family: "Microsoft YaHei", Arial, sans-serif;
                margin: 40px;
                color: #333;
            }}
            h1 {{
                color: #1a1a1a;
                border-bottom: 2px solid #366092;
                padding-bottom: 10px;
            }}
            h2 {{
                color: #366092;
                margin-top: 30px;
            }}
            .header {{
                margin-bottom: 30px;
            }}
            .info-grid {{
                display: grid;
                grid-template-columns: repeat(3, 1fr);
                gap: 15px;
                margin: 20px 0;
            }}
            .info-item {{
                background: #f5f5f5;
                padding: 15px;
                border-radius: 5px;
            }}
            .info-label {{
                font-size: 12px;
                color: #666;
            }}
            .info-value {{
                font-size: 18px;
                font-weight: bold;
                color: #333;
            }}
            .metrics-grid {{
                display: grid;
                grid-template-columns: repeat(4, 1fr);
                gap: 20px;
                margin: 20px 0;
            }}
            .metric-card {{
                background: #f8f9fa;
                border: 1px solid #ddd;
                border-radius: 8px;
                padding: 20px;
                text-align: center;
            }}
            .metric-value {{
                font-size: 24px;
                font-weight: bold;
                color: #366092;
            }}
            .metric-label {{
                font-size: 14px;
                color: #666;
                margin-top: 5px;
            }}
            .positive {{ color: #28a745; }}
            .negative {{ color: #dc3545; }}
            table {{
                width: 100%;
                border-collapse: collapse;
                margin-top: 20px;
            }}
            th, td {{
                border: 1px solid #ddd;
                padding: 10px;
                text-align: left;
            }}
            th {{
                background: #366092;
                color: white;
            }}
            tr:nth-child(even) {{
                background: #f9f9f9;
            }}
            .footer {{
                margin-top: 40px;
                text-align: center;
                color: #666;
                font-size: 12px;
            }}
        </style>
    </head>
    <body>
        <div class="header">
            <h1>投资组合报告</h1>
            <p><strong>组合名称：</strong>{portfolio_data.get('name', 'N/A')}</p>
            <p><strong>风险等级：</strong>{portfolio_data.get('risk_level', 'N/A')}</p>
            <p><strong>生成时间：</strong>{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}</p>
        </div>

        <h2>组合概览</h2>
        <div class="metrics-grid">
            <div class="metric-card">
                <div class="metric-value">¥{portfolio_data.get('total_value', 0):,.2f}</div>
                <div class="metric-label">总资产</div>
            </div>
            <div class="metric-card">
                <div class="metric-value">¥{portfolio_data.get('cash_balance', 0):,.2f}</div>
                <div class="metric-label">现金余额</div>
            </div>
            <div class="metric-card">
                <div class="metric-value {'positive' if portfolio_data.get('daily_change', 0) >= 0 else 'negative'}">
                    {portfolio_data.get('daily_change', 0):.2f}%
                </div>
                <div class="metric-label">日涨跌幅</div>
            </div>
            <div class="metric-card">
                <div class="metric-value {'positive' if portfolio_data.get('annual_return', 0) >= 0 else 'negative'}">
                    {portfolio_data.get('annual_return', 0):.2f}%
                </div>
                <div class="metric-label">年化收益</div>
            </div>
        </div>

        <h2>资产配置</h2>
        <div class="info-grid">
            <div class="info-item">
                <div class="info-label">初始资金</div>
                <div class="info-value">¥{portfolio_data.get('initial_balance', 0):,.2f}</div>
            </div>
            <div class="info-item">
                <div class="info-label">总市值</div>
                <div class="info-value">¥{portfolio_data.get('total_value', 0) - portfolio_data.get('cash_balance', 0):,.2f}</div>
            </div>
            <div class="info-item">
                <div class="info-label">持仓占比</div>
                <div class="info-value">{((portfolio_data.get('total_value', 0) - portfolio_data.get('cash_balance', 0)) / portfolio_data.get('total_value', 1) * 100):.2f}%</div>
            </div>
        </div>

        <h2>持仓明细</h2>
        {holdings_html if holdings_html else '<p>暂无持仓</p>'}

        <div class="footer">
            <p>本报告由股票分析与量化投资平台自动生成</p>
            <p>仅供参考，不构成投资建议</p>
        </div>
    </body>
    </html>
    '''
    return html


if __name__ == '__main__':
    # 测试代码
    print("导出工具模块测试")

    # 测试CSV导出
    test_data = [
        {'name': '张三', 'age': 25, 'city': '北京'},
        {'name': '李四', 'age': 30, 'city': '上海'},
        {'name': '王五', 'age': 35, 'city': '深圳'}
    ]

    try:
        csv_path = export_to_csv(test_data, 'test_export.csv')
        print(f"CSV导出成功: {csv_path}")
    except Exception as e:
        print(f"CSV导出失败: {e}")

    # 测试Excel导出
    try:
        excel_path = export_to_excel(test_data, 'test_export.xlsx', sheet_name='测试')
        print(f"Excel导出成功: {excel_path}")
    except ImportError as e:
        print(f"Excel导出跳过: {e}")
    except Exception as e:
        print(f"Excel导出失败: {e}")
